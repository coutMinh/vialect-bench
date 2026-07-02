from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl


DIALECT_ORDER = [
    "north_nonstandard",
    "thanh_hoa",
    "nghe_tinh",
    "binh_tri_thien",
    "south_central",
    "southern",
]

EXCEL_DIALECT_COLUMNS = {
    "Bắc Bộ": "north_nonstandard",
    "Thanh Hoá": "thanh_hoa",
    "Nghệ An": "nghe_tinh",
    "Hà Tĩnh - Quảng Bình - Quảng Trị": "binh_tri_thien",
    "Thừa Thiên Huế": "binh_tri_thien",
    "Quảng Nam": "south_central",
    "Quảng Ngãi": "south_central",
    "Phú Yên": "south_central",
    "Nam Bộ": "southern",
}

DATASET_PATHS = {
    "sentiment": Path("data/sentiment_100.jsonl"),
    "nli": Path("data/nli_100.jsonl"),
    "qa": Path("data/qa_100.jsonl"),
    "mcqa": Path("data/mcqa_100.jsonl"),
}


@dataclass(frozen=True)
class Rule:
    source: str
    target: str
    dialect: str
    category: str


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def split_alternatives(value: str) -> list[str]:
    cleaned = value.strip()
    if not cleaned:
        return []
    parts = re.split(r"\s*/\s*|\s*,\s*", cleaned)
    results = []
    for part in parts:
        part = re.sub(r"\s*\([^)]*\)", "", part).strip()
        if part:
            results.append(part)
    return results


def category_to_feature(category: str) -> str:
    lowered = category.casefold()
    if "interrogative" in lowered or "demonstrative" in lowered:
        return "demonstrative_interrogative"
    if "pronoun" in lowered or "kinship" in lowered:
        return "pronoun_address"
    if "negation" in lowered:
        return "negation_modality"
    if "particle" in lowered:
        return "discourse_particle"
    if "connective" in lowered or "aspect" in lowered:
        return "connective_aspect"
    if "template" in lowered or "idiom" in lowered:
        return "template_idiom"
    return "lexical_replacement"


def load_lexicon_rules(path: Path) -> list[Rule]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SystemExit(
            f"Expected an Excel lexicon (.xlsx/.xlsm), got '{path}'. "
            "Use --lexicon data/DIALECT_LEXICON_v2.xlsx."
        )

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    header = [
        str(cell).strip() if cell is not None else ""
        for cell in next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))
    ]
    standard_col = header.index("Tiếng chuẩn")
    category_col = header.index("Category")
    dialect_cols = {
        idx: EXCEL_DIALECT_COLUMNS[name]
        for idx, name in enumerate(header)
        if name in EXCEL_DIALECT_COLUMNS
    }

    rules: list[Rule] = []
    seen: set[tuple[str, str, str]] = set()
    for row in worksheet.iter_rows(min_row=3, values_only=True):
        if len(row) < len(header):
            row = tuple(row) + ("",) * (len(header) - len(row))

        source_value = row[standard_col]
        category_value = row[category_col]
        if source_value is None or category_value is None:
            continue

        source_forms = split_alternatives(str(source_value))
        if not source_forms:
            continue

        category = category_to_feature(str(category_value))
        source_set = {source.casefold() for source in source_forms}
        for col_idx, dialect in dialect_cols.items():
            target_value = row[col_idx]
            if target_value is None:
                continue

            target_forms = split_alternatives(str(target_value))
            if {target.casefold() for target in target_forms} == source_set:
                continue

            for source in source_forms:
                for target in target_forms:
                    if source.casefold() == target.casefold():
                        continue
                    dedup_key = (dialect, source.casefold(), target.casefold())
                    if dedup_key in seen:
                        continue
                    seen.add(dedup_key)
                    rules.append(
                        Rule(
                            source=source,
                            target=target,
                            dialect=dialect,
                            category=category,
                        )
                    )

    workbook.close()
    return rules


def regex_for_source(source: str) -> re.Pattern[str]:
    return re.compile(rf"(?<!\w){re.escape(source)}(?!\w)", flags=re.IGNORECASE)


def preserve_case(replacement: str, matched: str) -> str:
    if matched.isupper():
        return replacement.upper()
    if matched[:1].isupper():
        return replacement[:1].upper() + replacement[1:]
    return replacement


def apply_rules(text: str, dialect: str, rules: list[Rule]) -> tuple[str, list[dict[str, str]]]:
    changed = text
    fired: list[dict[str, str]] = []
    dialect_rules = sorted(
        [rule for rule in rules if rule.dialect == dialect],
        key=lambda rule: len(rule.source),
        reverse=True,
    )
    for rule in dialect_rules:
        pattern = regex_for_source(rule.source)

        def replace(match: re.Match[str]) -> str:
            if rule.category == "pronoun_address" and match.group(0)[:1].isupper():
                return match.group(0)
            target = preserve_case(rule.target, match.group(0))
            fired.append(
                {
                    "source": match.group(0),
                    "target": target,
                    "feature": rule.category,
                }
            )
            return target

        changed = pattern.sub(replace, changed)
    return changed, fired


def build_sentiment_case(row: dict[str, Any], index: int, rules: list[Rule]) -> dict[str, Any]:
    standard_text = row["Sentence"]
    case = {
        "id": f"sentiment_{index:03d}",
        "task": "sentiment",
        "label": row["Emotion"],
        "source_dataset": row["_source_dataset"],
        "source_index": row["_source_index"],
        "note": "Rule-based dialect candidates for probing; require native-speaker review before reporting.",
        "variants": {"standard": {"text": standard_text}},
    }
    add_text_variants(case, standard_text, rules, field_name="text")
    return case


def build_nli_case(row: dict[str, Any], index: int, rules: list[Rule]) -> dict[str, Any]:
    hypothesis = row["hypothesis"]
    case = {
        "id": f"nli_{index:03d}",
        "task": "nli",
        "label": row["label"],
        "source_dataset": row["_source_dataset"],
        "source_index": row["_source_index"],
        "uid": row.get("uid"),
        "note": "Rule-based dialect candidates paraphrase only the hypothesis; premise stays standard.",
        "variants": {
            "standard": {
                "premise": row["premise"],
                "hypothesis": hypothesis,
            }
        },
    }
    add_text_variants(case, hypothesis, rules, field_name="hypothesis")
    for variant_name, variant in case["variants"].items():
        if variant_name != "standard":
            variant["premise"] = row["premise"]
    return case


def build_qa_case(row: dict[str, Any], index: int, rules: list[Rule]) -> dict[str, Any]:
    question = row["question"]
    answers = row.get("answers") or {}
    answer_text = answers.get("text") if isinstance(answers, dict) else answers
    case = {
        "id": f"qa_{index:03d}",
        "task": "qa",
        "answers": answer_text,
        "source_dataset": row["_source_dataset"],
        "source_index": row["_source_index"],
        "source_id": row.get("id"),
        "source_title": row.get("title"),
        "note": "Rule-based dialect candidates paraphrase only the question; context and answer span stay standard.",
        "variants": {
            "standard": {
                "context": row["context"],
                "question": question,
            }
        },
    }
    add_text_variants(case, question, rules, field_name="question")
    return case


def build_mcqa_case(row: dict[str, Any], index: int, rules: list[Rule]) -> dict[str, Any]:
    question = row["question"]
    case = {
        "id": f"mcqa_{index:03d}",
        "task": "mcqa",
        "label": row["answer"],
        "source_dataset": row["_source_dataset"],
        "source_index": row["_source_index"],
        "source_title": row.get("title"),
        "note": "Rule-based dialect candidates paraphrase only the question; article/options/answer stay standard.",
        "variants": {
            "standard": {
                "context": row["article"],
                "question": question,
                "options": row["options"],
            }
        },
    }
    add_text_variants(case, question, rules, field_name="question")
    return case


def add_text_variants(case: dict[str, Any], text: str, rules: list[Rule], field_name: str) -> None:
    for dialect in DIALECT_ORDER:
        paraphrase, _ = apply_rules(text, dialect, rules)
        case["variants"][dialect] = {field_name: paraphrase}


def build_cases(data_dir: Path, rules: list[Rule], limit: int | None) -> list[dict[str, Any]]:
    builders = {
        "sentiment": build_sentiment_case,
        "nli": build_nli_case,
        "qa": build_qa_case,
        "mcqa": build_mcqa_case,
    }
    cases: list[dict[str, Any]] = []
    for task, relative_path in DATASET_PATHS.items():
        rows = load_jsonl(data_dir / relative_path.name)
        if limit is not None:
            rows = rows[:limit]
        for index, row in enumerate(rows, start=1):
            cases.append(builders[task](row, index, rules))
    return cases


def summarize(cases: list[dict[str, Any]]) -> dict[str, Any]:
    summary: dict[str, Any] = {"total_cases": len(cases), "tasks": {}, "dialects": {}}
    for case in cases:
        task = case["task"]
        summary["tasks"].setdefault(task, 0)
        summary["tasks"][task] += 1
        for dialect in case["variants"]:
            if dialect == "standard":
                continue
            dialect_summary = summary["dialects"].setdefault(dialect, {"variants": 0})
            dialect_summary["variants"] += 1
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create rule-based dialect probe cases from the 100-sample task datasets."
    )
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--lexicon", type=Path, default=Path("data/DIALECT_LEXICON_v2.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("examples/probe_cases_rule_based.jsonl"))
    parser.add_argument("--summary", type=Path, default=Path("outputs/rule_based_probe_summary.json"))
    parser.add_argument("--limit-per-task", type=int)
    args = parser.parse_args()

    rules = load_lexicon_rules(args.lexicon)
    cases = build_cases(args.data_dir, rules, args.limit_per_task)
    write_jsonl(args.output, cases)
    args.summary.parent.mkdir(parents=True, exist_ok=True)
    args.summary.write_text(
        json.dumps(summarize(cases), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
